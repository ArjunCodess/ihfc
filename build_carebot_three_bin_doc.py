from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output" / "documents"
QA_DIR = ROOT / "output" / "documents" / "carebot_three_bin_qa"
DOCX_PATH = OUT_DIR / "CAREBOT-THREE-BIN-DESIGN.docx"
DIAGRAM_PATH = QA_DIR / "three_bin_layout.png"

OUT_DIR.mkdir(parents=True, exist_ok=True)
QA_DIR.mkdir(parents=True, exist_ok=True)

NAVY = "193F50"
BLUE = "DDEBF3"
TEAL = "DDEFEA"
PALE = "F5F7F8"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def font(size):
    try:
        return ImageFont.truetype("C:/Windows/Fonts/arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def build_diagram():
    image = Image.new("RGB", (1400, 620), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(42)
    body_font = font(28)
    small_font = font(23)
    draw.text((60, 35), "CareBot three-bin layout", fill="#193F50", font=title_font)
    draw.text((60, 95), "Top view, front of robot at the bottom", fill="#52636C", font=small_font)

    x0, y0, w, h = 100, 175, 1200, 350
    draw.rounded_rectangle((x0, y0, x0 + w, y0 + h), radius=24, outline="#193F50", width=5, fill="#F5F7F8")
    bins = [
        (150, 220, 390, 455, "Small bin A", "2 cubes", "Servo flap 1", "#DDEFEA"),
        (430, 205, 970, 455, "Central bin", "6 cubes in a low 3 x 2 layout", "Servo flap 2", "#DDEBF3"),
        (1010, 220, 1250, 455, "Small bin B", "2 cubes", "Servo flap 3", "#DDEFEA"),
    ]
    for left, top, right, bottom, name, load, flap, fill in bins:
        draw.rounded_rectangle((left, top, right, bottom), radius=18, outline="#193F50", width=4, fill=fill)
        cx = (left + right) // 2
        draw.text((cx, top + 45), name, fill="#193F50", font=body_font, anchor="ma")
        draw.text((cx, top + 100), load, fill="#000000", font=small_font, anchor="ma")
        draw.line((left + 24, bottom - 55, right - 24, bottom - 55), fill="#193F50", width=4)
        draw.text((cx, bottom - 32), flap, fill="#000000", font=small_font, anchor="mm")
    draw.polygon([(680, 585), (655, 545), (705, 545)], fill="#193F50")
    draw.text((680, 605), "Direction of travel", fill="#193F50", font=body_font, anchor="ms")
    image.save(DIAGRAM_PATH)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_table_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, widths)
    for idx, heading in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = heading
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            set_run_font(run, size=9.5, bold=True, color=RGBColor(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[idx])
            if row_index % 2:
                set_cell_shading(cells[idx], PALE)
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                set_run_font(run, size=9.5)
    doc.add_paragraph()
    return table


def add_body(doc, text, bold_start=None):
    p = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        r1 = p.add_run(bold_start)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_start):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_run_font(p.add_run(text))
    return p


def add_number(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.20)
    lead = p.add_run(f"{number}. ")
    set_run_font(lead)
    r1 = p.add_run(title + " ")
    set_run_font(r1, bold=True)
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_run_font(run, bold=True, color=BLACK)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=RGBColor(82, 99, 108))
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


build_diagram()
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.66)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.8)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

title_style = doc.styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(27)
title_style.font.bold = True
title_style.font.color.rgb = BLACK
title_style.paragraph_format.space_after = Pt(10)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size, before, after in [("Heading 1", 17, 15, 7), ("Heading 2", 13, 11, 5)]:
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph(style="Title")
title.add_run("CareBot three bin patient delivery design")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
run = subtitle.add_run("Separate design document for the first robot")
set_run_font(run, size=13, color=RGBColor(82, 99, 108))

add_body(doc, "CareBot is a two-wheel robotic car that carries ten preloaded Red Cross patient cubes in three bins. Two small bins carry two cubes each, and the larger central bin carries six. A servo-controlled bottom flap releases each group at its assigned field location. An ESP32-CAM can identify coloured cylindrical patients and select a stored destination route, while the Uno R3 controls movement and all three bin flaps.")

add_heading(doc, "Design decision", 1)
add_body(doc, "This version removes the pickup lift, colour sensor, sorting flap, and seven storage boxes from the earlier CareBot concept. The cubes are loaded before the run, so the robot only has to navigate and open the correct flap. That makes the mechanism lighter, easier to build, and easier to debug.")

picture = doc.add_paragraph()
picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
picture.paragraph_format.keep_with_next = True
picture.add_run().add_picture(str(DIAGRAM_PATH), width=Inches(6.55))
caption = doc.add_paragraph("Figure 1  Proposed three-bin arrangement")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(10)
for run in caption.runs:
    set_run_font(run, size=8.5, color=RGBColor(82, 99, 108))

add_heading(doc, "Mechanical design", 1)
add_table(
    doc,
    ["Part", "Design and job"],
    [
        ["Robot base", "Two geared drive motors, two wheels, one caster, and a rigid upper deck."],
        ["Small bin A", "Holds two preloaded cubes and releases them through its own bottom flap."],
        ["Central bin", "Holds six cubes in a low arrangement and releases all six through one wide bottom flap."],
        ["Small bin B", "Holds the final two cubes and releases them through its own bottom flap."],
        ["Flap servos", "Use one servo per bin. Use an MG90S for the six-cube flap; SG90 servos may be enough for the small bins after a load test."],
    ],
    [1.35, 5.0],
)

add_body(doc, "Keep the three bins centred between the drive wheels so the robot does not pull to one side as a bin empties. The preferred central-bin layout is one low layer of six cubes rather than a tall stack, because a lower centre of gravity reduces tipping during the 45 degree turn.")
add_body(doc, "Let C be the measured cube side. A small bin needs an internal floor close to 2C by C, plus 4 to 6 mm of free clearance. The central bin needs roughly 3C by 2C, again with clearance. Build a cardboard version around the real cubes before cutting acrylic or plywood.")

add_heading(doc, "Proposed robot envelope", 2)
add_table(
    doc,
    ["Dimension", "Proposed maximum", "Reason"],
    [
        ["Width", "230 mm", "Fits beside SealBot inside the previously planned 480 mm start-zone width."],
        ["Length", "240 mm", "Leaves clearance inside the 280 mm start-zone length."],
        ["Height", "280 mm", "Allows the camera and bin walls while retaining the earlier CareBot envelope."],
    ],
    [1.0, 1.3, 4.05],
)
add_body(doc, "These are outer limits, not final bin measurements. Confirm the finished size with the actual cubes and the competition measuring box.")

add_heading(doc, "Operating sequence", 1)
add_number(doc, 1, "Load the bins.", "Place two cubes in small bin A, six in the central bin, and two in small bin B before the run.")
add_number(doc, 2, "Travel along the main path.", "The Uno follows the line or stored wheel-encoder distance until it detects the first release marker.")
add_number(doc, 3, "Release the first pair.", "The first small-bin servo opens its bottom flap. Both cubes fall into the first target area, then the flap closes.")
add_number(doc, 4, "Take the left branch.", "The robot turns left and stops turning when the line sensor finds the 45 degree branch. Encoder calibration can be used if the field has no branch line.")
add_number(doc, 5, "Reach the central target.", "The robot follows the second path for the calibrated distance. Replace the phrase ten steps with an encoder count or field marker, because software loop counts change with battery voltage and wheel slip.")
add_number(doc, 6, "Release six cubes.", "The central MG90S servo opens the large flap and drops all six cubes. The robot waits briefly, then closes the flap.")
add_number(doc, 7, "Reach the final target.", "The robot continues to the last marker and opens small bin B, releasing the final two cubes.")

doc.add_page_break()
add_heading(doc, "Electronics and control", 1)
add_table(
    doc,
    ["Part", "Quantity", "Job"],
    [
        ["Uno R3", "1", "Runs the movement sequence, reads the line sensor and encoders, and commands the three bin servos."],
        ["ESP32-CAM", "1", "Recognises coloured cylindrical patients and sends a short patient label to the Uno."],
        ["DRV8833", "1", "Controls the left and right drive motors."],
        ["Three-channel line sensor", "1", "Follows the route and detects field branches or release markers."],
        ["SG90 servo", "2", "Opens the two small-bin flaps after load testing."],
        ["MG90S servo", "1", "Opens the wider flap carrying six cubes."],
        ["Wheel encoders", "2", "Measure travel distance and improve turn repeatability."],
        ["2S battery and regulators", "1 set", "Supplies the motors, servos, Uno, and camera through separate regulated rails."],
    ],
    [1.55, 0.75, 4.05],
)

add_heading(doc, "Why the PCA9685 is optional", 2)
add_body(doc, "The revised robot has only three servos, so the Uno can generate all three servo signals directly. A PCA9685 adds wiring and another board without solving a current problem. Keep it only if the design later grows beyond the Uno pin budget or repeated servo timing affects movement.")

add_heading(doc, "Power arrangement", 2)
add_bullet(doc, "Feed the Uno and ESP32-CAM from a regulated 5 V logic rail.")
add_bullet(doc, "Feed the servos from a separate 5 V actuator rail with enough current for a flap starting under load.")
add_bullet(doc, "Connect the grounds of the logic rail, actuator rail, motor driver, Uno, and camera.")
add_bullet(doc, "Do not power the three servos through the Uno 5 V pin.")

add_heading(doc, "Vision and route mapping", 1)
add_body(doc, "The ESP32-CAM handles image capture and classification. Train the model with the actual coloured cylindrical patients under several field lighting conditions, viewing angles, and distances. The camera should send only a compact result such as RED, GREEN, YELLOW, or UNKNOWN to the Uno. The Uno then chooses the stored route and controls the motors and flap servo. If no target is visible, the robot keeps every flap closed and returns to its search routine or last marker.")
add_body(doc, "For this fixed field, predefined coordinates should mean stored route instructions, marker counts, and encoder distances. The camera alone does not provide a reliable global X and Y position. True coordinate navigation would require field markers, wheel odometry with correction, or another localisation system.")
add_bullet(doc, "When the colour is known, the Uno selects the stored route and opens only the bin or carrier assigned to that destination.")
add_bullet(doc, "When the result is unknown, the robot stops, adjusts its viewing position, classifies again, and keeps every flap closed.")

add_heading(doc, "Scope of the camera system", 2)
add_body(doc, "The three-bin mechanism completes the preloaded cube-delivery task. The camera can recognise coloured cylindrical patients and decide where they belong, but the current mechanical design does not pick them up. If the rules require the same robot to collect and carry those cylinders, add a separate front scoop or gripper and a dedicated carrier. Do not claim autonomous placement of loose cylindrical patients until that mechanism has been built and tested.")

add_heading(doc, "Firmware sequence", 1)
add_table(
    doc,
    ["State", "Exit condition", "Next action"],
    [
        ["Start", "Start button pressed", "Close all flaps and begin the main path."],
        ["Main path", "First target marker detected", "Stop and open small bin A."],
        ["Left turn", "Branch line reacquired", "Follow the secondary path."],
        ["Central path", "Central marker or encoder count reached", "Stop and open the central bin."],
        ["Final path", "Final marker detected", "Stop and open small bin B."],
        ["Finish", "Release delay completed", "Close the flap, stop the motors, and ignore further commands."],
    ],
    [1.25, 2.55, 2.55],
)

add_heading(doc, "Build and test order", 1)
add_number(doc, 1, "Test one flap.", "Build one cardboard bin, load it with the correct number of cubes, and confirm that the flap opens without the servo stalling.")
add_number(doc, 2, "Build all three bins.", "Check that every cube leaves cleanly and that no cube catches on a hinge, servo horn, or wall.")
add_number(doc, 3, "Calibrate movement.", "Tune straight driving, line following, the left turn, and each stop position before adding camera code.")
add_number(doc, 4, "Add the camera.", "Train the colour model and test its label output separately, then connect its serial result to the Uno route table.")
add_number(doc, 5, "Run the full load.", "Repeat the complete ten-cube mission with a fully charged and partly discharged battery. Adjust encoder counts and release delays only after several runs.")

doc.add_page_break()
add_heading(doc, "Presentation script", 1)
script_paragraphs = [
    "I will explain the basic design and working of our first robot, CareBot. It is a robotic car that transports patient cubes and releases each group at its assigned location. The movement is automatic, and the robot uses separate servo-controlled bins for placement.",
    "At the beginning of the task, the rules allow us to place the Red Cross marked patient cubes directly on the robot. We will load all ten cubes before the run. This removes the need for a pickup or sorting mechanism during the cube-delivery sequence.",
    "CareBot has three storage bins. The two smaller bins hold two cubes each, while the larger central bin holds six cubes. Every bin has a flap at the bottom, and each flap has its own servo motor. The program releases a group by moving only the servo connected to that bin.",
    "The robot first moves straight along the main path. When the first target marker is detected, the robot stops and opens the first small bin. The two cubes fall into the designated area, the flap closes, and the robot begins the next movement.",
    "CareBot then turns left onto the 45 degree branch. Instead of relying on a rough software step count, the robot follows the branch line and uses wheel-encoder counts or a field marker to measure the next section. At the central target, it opens the large bin and releases all six cubes.",
    "After the central release, the robot continues to the final point. When the last marker is detected, the servo under the second small bin opens its flap and drops the remaining two cubes. The robot then closes the flap and stops.",
    "An Uno R3 controls the drive motors, line sensor, wheel encoders, and the three flap servos. This version has only three servos, so the Uno can drive them directly and a PCA9685 is optional. The servos use a separate regulated power rail rather than drawing their current through the Uno.",
    "We also plan to use an ESP32-CAM for the vision task. We will train a small image-classification model with the actual coloured cylindrical patients under different field lighting and viewing angles. The camera will identify the patient's colour and send a short label to the Uno.",
    "The Uno will match that colour label to a stored route made from line markers and encoder distances. These stored routes are more reliable for our fixed field than claiming that the camera gives an exact global coordinate. If the camera is unsure, the robot will stop and classify the patient again instead of opening a flap.",
    "The three-bin mechanism completes the preloaded cube-delivery task. If CareBot must also collect loose cylindrical patients, we will add a front scoop or gripper and a separate carrier. The camera decides what the object is and where it belongs, while a mechanical pickup is still required to move it.",
]
for paragraph in script_paragraphs:
    add_body(doc, paragraph)

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.text = "CareBot three bin patient delivery design  |  7 September 2026"
    for run in footer.runs:
        set_run_font(run, size=8.5, color=RGBColor(82, 99, 108))
    add_page_number(section.footer.add_paragraph())

doc.core_properties.title = "CareBot three bin patient delivery design"
doc.core_properties.subject = "Mechanical design, movement sequence, electronics, vision, and build plan"
doc.core_properties.author = "CareBot project team"
doc.core_properties.keywords = "CareBot, Arduino Uno, ESP32-CAM, patient cubes, servo bins"
doc.save(DOCX_PATH)
print(DOCX_PATH)
